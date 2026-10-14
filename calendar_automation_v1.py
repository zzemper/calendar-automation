# -*- coding: utf-8 -*-
"""
Created on Tue Oct 13 14:33:09 2026

@author: 3354745
"""

import io
import re
from datetime import date, datetime, time, timedelta

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# PAGE SETUP
# ============================================================

st.set_page_config(
    page_title="Holiday DF Shipping Calendars Automation",
    page_icon="📅",
    layout="wide",
)

st.title("📅 Holiday DF Shipping Calendars Automation")
st.write(
    "Select fiscal weeks and calendars, then click each date to cycle "
    "between Not selected, Open, and Closed."
)


# ============================================================
# CONSTANTS
# ============================================================

BASE_FISCAL_WEEK = 1
BASE_FISCAL_WEEK_START = date(2026, 2, 1)  # Sunday
MAX_FISCAL_WEEK = 52

AVAILABLE_CALENDARS = [
    "SHIP_DF_860_TG",
    "New calendar",
]

MINUTE_OPTIONS = [
    "First minute of the day",
    "Last minute of the day",
    "2nd to last minute of the day",
]

DAY_STATUS_SEQUENCE = ["Not selected", "Open", "Closed"]

STATUS_STYLE = {
    "Not selected": {
        "icon": "⚪",
        "label": "\nNot selected",
    },
    "Open": {
        "icon": "🟢",
        "label": "\nOPEN",
    },
    "Closed": {
        "icon": "🔴",
        "label": "\nCLOSED",
    },
}

WEEKDAY_NAMES = [
    "Sunday",
    "Monday",
    "Tuesday",
    "Wednesday",
    "Thursday",
    "Friday",
    "Saturday",
]

OUTPUT_COLUMNS = [
    "*Cal",
    "*Eff",
    "AllocWgt",
    "AltCal",
    "Avail",
    "Descr",
    "Opt",
    "PerWgt",
    "Repeat",
]


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def safe_key(value):
    """Convert text into a safe Streamlit widget/session key."""
    return re.sub(r"[^A-Za-z0-9]+", "_", value).strip("_")


def fiscal_week_start(week_number):
    """Return the Sunday starting the requested 2026 fiscal week."""
    return BASE_FISCAL_WEEK_START + timedelta(
        weeks=week_number - BASE_FISCAL_WEEK
    )


def dates_for_fiscal_week(week_number):
    """Return the seven dates in a fiscal week, Sunday through Saturday."""
    start_date = fiscal_week_start(week_number)
    return [start_date + timedelta(days=offset) for offset in range(7)]


def date_status_key(calendar_name, calendar_date):
    """Create a unique state key for a calendar and date."""
    return (
        f"day_status__{safe_key(calendar_name)}__"
        f"{calendar_date.isoformat()}"
    )


def get_day_status(calendar_name, calendar_date):
    """Return a date's current state, defaulting to gray/not selected."""
    key = date_status_key(calendar_name, calendar_date)
    return st.session_state.get(key, "Not selected")


def cycle_day_status(calendar_name, calendar_date):
    """Cycle gray -> green/open -> red/closed -> gray."""
    key = date_status_key(calendar_name, calendar_date)
    current_status = st.session_state.get(key, "Not selected")
    current_index = DAY_STATUS_SEQUENCE.index(current_status)
    next_index = (current_index + 1) % len(DAY_STATUS_SEQUENCE)
    st.session_state[key] = DAY_STATUS_SEQUENCE[next_index]


def clear_calendar_dates(calendar_name, fiscal_weeks):
    """Reset all visible dates for one calendar to gray/not selected."""
    for week_number in fiscal_weeks:
        for calendar_date in dates_for_fiscal_week(week_number):
            key = date_status_key(calendar_name, calendar_date)
            st.session_state[key] = "Not selected"


def mark_all_calendar_dates(calendar_name, fiscal_weeks, status):
    """Mark all visible dates for one calendar Open or Closed."""
    for week_number in fiscal_weeks:
        for calendar_date in dates_for_fiscal_week(week_number):
            key = date_status_key(calendar_name, calendar_date)
            st.session_state[key] = status


def build_effective_datetimes(calendar_date, minute_option):
    """Return the two *Eff values required for one selected date."""
    if minute_option == "First minute of the day":
        first_record = datetime.combine(calendar_date, time(0, 0))
        second_record = datetime.combine(calendar_date, time(0, 1))

    elif minute_option == "Last minute of the day":
        first_record = datetime.combine(calendar_date, time(23, 59))
        second_record = datetime.combine(
            calendar_date + timedelta(days=1),
            time(0, 0),
        )

    else:  # 2nd to last minute of the day
        first_record = datetime.combine(calendar_date, time(23, 58))
        second_record = datetime.combine(calendar_date, time(23, 59))

    return first_record, second_record


def build_output_dataframe(
    selected_calendars,
    fiscal_weeks,
    minute_selection_by_calendar,
    description,
    add_repeat,
):
    """
    Build the final import-format dataframe.

    Export rules:
    - Gray/Not selected dates are excluded.
    - Open dates create two rows with Opt = 2.
    - Closed dates create two rows with Opt = 1.
    - If add_repeat is enabled, one additional row is created
      for the day immediately after the latest selected date.
    - The additional repeat row receives Opt = 5 and Repeat = 7D.
    """

    rows = []

    for calendar_name in selected_calendars:
        minute_option = minute_selection_by_calendar[
            calendar_name
        ]

        calendar_rows = []
        selected_calendar_dates = []

        for week_number in fiscal_weeks:
            for calendar_date in dates_for_fiscal_week(
                week_number
            ):
                status = get_day_status(
                    calendar_name,
                    calendar_date
                )

                # Gray dates are intentionally excluded.
                if status == "Not selected":
                    continue

                # Track the actual calendar date selected by the user.
                selected_calendar_dates.append(
                    calendar_date
                )

                if status == "Open":
                    opt_value = 2
                else:
                    opt_value = 1

                first_eff, second_eff = (
                    build_effective_datetimes(
                        calendar_date,
                        minute_option
                    )
                )

                # First record for the selected date.
                calendar_rows.append(
                    {
                        "*Cal": calendar_name,
                        "*Eff": first_eff,
                        "AllocWgt": 0,
                        "AltCal": "",
                        "Avail": 0,
                        "Descr": description.strip(),
                        "Opt": opt_value,
                        "PerWgt": 0,
                        "Repeat": "",
                    }
                )

                # Second record for the selected date.
                calendar_rows.append(
                    {
                        "*Cal": calendar_name,
                        "*Eff": second_eff,
                        "AllocWgt": 0,
                        "AltCal": "",
                        "Avail": 0,
                        "Descr": description.strip(),
                        "Opt": opt_value,
                        "PerWgt": 0,
                        "Repeat": "",
                    }
                )

        # ----------------------------------------------------
        # ADD THE NEW REPEAT ROW
        # ----------------------------------------------------

        if (
            add_repeat
            and selected_calendar_dates
        ):
            latest_selected_date = max(
                selected_calendar_dates
            )

            repeat_date = (
                latest_selected_date
                + timedelta(days=1)
            )

            if minute_option == "First minute of the day":
                repeat_effective_datetime = (
                    datetime.combine(
                        repeat_date,
                        time(0, 0)
                    )
                )

            elif minute_option == "Last minute of the day":
                repeat_effective_datetime = (
                    datetime.combine(
                        repeat_date,
                        time(23, 59)
                    )
                )

            else:
                # 2nd to last minute of the day
                repeat_effective_datetime = (
                    datetime.combine(
                        repeat_date,
                        time(23, 58)
                    )
                )

            calendar_rows.append(
                {
                    "*Cal": calendar_name,
                    "*Eff": repeat_effective_datetime,
                    "AllocWgt": 0,
                    "AltCal": "",
                    "Avail": 0,
                    "Descr": description.strip(),
                    "Opt": 5,
                    "PerWgt": 0,
                    "Repeat": "7D",
                }
            )

        calendar_rows.sort(
            key=lambda row: row["*Eff"]
        )

        rows.extend(calendar_rows)

    output_df = pd.DataFrame(
        rows,
        columns=OUTPUT_COLUMNS
    )

    if not output_df.empty:
        output_df["*Eff"] = pd.to_datetime(
            output_df["*Eff"]
        )

        output_df = output_df.sort_values(
            by=["*Cal", "*Eff"],
            kind="stable"
        ).reset_index(drop=True)

    return output_df


def create_excel_file(output_df):
    """Create the formatted Excel import file in memory."""
    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
        datetime_format="m/d/yyyy h:mm",
    ) as writer:
        output_df.to_excel(
            writer,
            index=False,
            sheet_name="Calendar Import",
        )

        worksheet = writer.book["Calendar Import"]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        header_font = Font(
            color="FFFFFF",
            bold=True,
        )

        for header_cell in worksheet[1]:
            header_cell.fill = header_fill
            header_cell.font = header_font
            header_cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for date_cell in worksheet["B"][1:]:
            date_cell.number_format = "m/d/yyyy h:mm"

        widths = {
            "A": 26,
            "B": 20,
            "C": 12,
            "D": 12,
            "E": 10,
            "F": 30,
            "G": 8,
            "H": 12,
            "I": 10,
        }

        for column_letter, width in widths.items():
            worksheet.column_dimensions[column_letter].width = width

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center")

    output.seek(0)
    return output


def selected_date_summary(calendar_name, fiscal_weeks):
    """Count the selected open, closed, and gray dates."""
    counts = {
        "Open": 0,
        "Closed": 0,
        "Not selected": 0,
    }

    for week_number in fiscal_weeks:
        for calendar_date in dates_for_fiscal_week(week_number):
            status = get_day_status(calendar_name, calendar_date)
            counts[status] += 1

    return counts


# ============================================================
# STEP 1: FISCAL WEEKS
# ============================================================

st.header("Step 1: Select fiscal weeks")

st.info(
    "Fiscal week 40 begins Sunday, November 1, 2026. "
    "Each following fiscal week begins seven days later."
)

week_col1, week_col2 = st.columns(2)

with week_col1:
    starting_week = st.number_input(
        "Starting fiscal week",
        min_value=BASE_FISCAL_WEEK,
        max_value=MAX_FISCAL_WEEK,
        value=BASE_FISCAL_WEEK,
        step=1,
    )

with week_col2:
    ending_week = st.number_input(
        "Ending fiscal week",
        min_value=BASE_FISCAL_WEEK,
        max_value=MAX_FISCAL_WEEK,
        value=47,
        step=1,
    )

if ending_week < starting_week:
    st.error("Ending fiscal week must be the same as or later than the starting week.")
    st.stop()

fiscal_weeks = list(range(int(starting_week), int(ending_week) + 1))

week_reference = pd.DataFrame(
    [
        {
            "Fiscal Week": week_number,
            "Start Date": fiscal_week_start(week_number),
            "End Date": fiscal_week_start(week_number) + timedelta(days=6),
        }
        for week_number in fiscal_weeks
    ]
)

st.dataframe(
    week_reference,
    hide_index=True,
    use_container_width=True,
)


# ============================================================
# STEP 2: CALENDAR SELECTION
# ============================================================

st.header("Step 2: Select calendars to update")

existing_calendars = st.multiselect(
    "Existing calendars",
    options=[
        "SHIP_DF_860_TG"
    ],
    default=["SHIP_DF_860_TG"]
)

st.markdown(
    """
    ### Additional New Calendars

    Paste one calendar name per line.

    These names will be written directly into the *Cal column
    of the final export file.
    """
)

new_calendar_text = st.text_area(
    "New calendar names",
    height=150,
    placeholder=(
        "SHIP_DF_999_TG\n"
        "SHIP_DF_998_TG\n"
        "SHIP_DF_997_TG"
    )
)

new_calendars = []

for row in new_calendar_text.splitlines():

    cleaned_name = row.strip()

    if cleaned_name:
        new_calendars.append(
            cleaned_name
        )

selected_calendars = (
    existing_calendars
    + new_calendars
)

# Remove accidental duplicates.
selected_calendars = list(
    dict.fromkeys(selected_calendars)
)

if not selected_calendars:

    st.warning(
        "Select or enter at least one calendar."
    )

    st.stop()


# ============================================================
# STEP 3: MINUTE SELECTION
# ============================================================

st.header("Step 3: Select the minutes for each calendar")

minute_selection_by_calendar = {}

for calendar_name in selected_calendars:
    calendar_col, minute_col = st.columns([2, 3])

    with calendar_col:
        st.markdown(f"**{calendar_name}**")

    with minute_col:
        minute_selection_by_calendar[calendar_name] = st.selectbox(
            f"Minutes for {calendar_name}",
            options=MINUTE_OPTIONS,
            key=f"minute_option__{safe_key(calendar_name)}",
            label_visibility="collapsed",
        )


# ============================================================
# STEP 4: CLICKABLE CALENDARS
# ============================================================

st.header("Step 4: Mark calendar dates")

st.info(
    "Each date starts gray and is excluded from export. Click once to mark it "
    "green/Open, click again to mark it red/Closed, and click a third time to "
    "return it to gray/Not selected. Include seven normally scheduled days at "
    "the end when a 7D automatic repeat should be added."
)

st.caption("Legend: ⚪ Not selected   |   🟢 Open   |   🔴 Closed")

for calendar_name in selected_calendars:
    st.subheader(calendar_name)
    st.caption(
        f"Minute rule: {minute_selection_by_calendar[calendar_name]}"
    )

    action_col1, action_col2, action_col3, spacer = st.columns([1, 1, 1, 3])

    with action_col1:
        if st.button(
            "Mark all Open",
            key=f"all_open__{safe_key(calendar_name)}",
            use_container_width=True,
        ):
            mark_all_calendar_dates(calendar_name, fiscal_weeks, "Open")
            st.rerun()

    with action_col2:
        if st.button(
            "Mark all Closed",
            key=f"all_closed__{safe_key(calendar_name)}",
            use_container_width=True,
        ):
            mark_all_calendar_dates(calendar_name, fiscal_weeks, "Closed")
            st.rerun()

    with action_col3:
        if st.button(
            "Clear all",
            key=f"clear_all__{safe_key(calendar_name)}",
            use_container_width=True,
        ):
            clear_calendar_dates(calendar_name, fiscal_weeks)
            st.rerun()

    header_columns = st.columns([0.7, 1, 1, 1, 1, 1, 1, 1])
    header_columns[0].markdown("**Fiscal Week**")

    for index, weekday_name in enumerate(WEEKDAY_NAMES, start=1):
        header_columns[index].markdown(
            f"<div style='text-align:center'><b>{weekday_name}</b></div>",
            unsafe_allow_html=True,
        )

    for week_number in fiscal_weeks:
        row_columns = st.columns([0.7, 1, 1, 1, 1, 1, 1, 1])
        row_columns[0].markdown(f"### FW {week_number}")

        week_dates = dates_for_fiscal_week(week_number)

        for day_index, calendar_date in enumerate(week_dates, start=1):
            status = get_day_status(calendar_name, calendar_date)
            style = STATUS_STYLE[status]
            button_label = (
                f"{style['icon']} {calendar_date.strftime('%b %d')}\n"
                f"{style['label']}"
            )

            with row_columns[day_index]:
                if st.button(
                    button_label,
                    key=(
                        f"day_button__{safe_key(calendar_name)}__"
                        f"{calendar_date.isoformat()}"
                    ),
                    use_container_width=True,
                    help=(
                        "Click to cycle: Not selected → Open → Closed → Not selected"
                    ),
                ):
                    cycle_day_status(calendar_name, calendar_date)
                    st.rerun()

    counts = selected_date_summary(calendar_name, fiscal_weeks)
    summary_col1, summary_col2, summary_col3 = st.columns(3)
    summary_col1.metric("Open dates", counts["Open"])
    summary_col2.metric("Closed dates", counts["Closed"])
    summary_col3.metric("Not selected", counts["Not selected"])

    st.divider()


# ============================================================
# STEP 5: DESCRIPTION, PREVIEW, AND DOWNLOAD
# ============================================================

st.header("Step 5: Describe and download the file")

description = st.text_input(
    "Description for the Descr column",
    value="turkeys 2026",
)

add_repeat = st.checkbox(
    "Add the Opt 5 / 7D repeat row for each calendar",
    value=True,
    help=(
        "Adds one new row for the day immediately after the "
        "latest selected date. The row receives Opt = 5 and "
        "Repeat = 7D."
    ),
)

output_df = build_output_dataframe(
    selected_calendars=selected_calendars,
    fiscal_weeks=fiscal_weeks,
    minute_selection_by_calendar=minute_selection_by_calendar,
    description=description,
    add_repeat=add_repeat,
)

if output_df.empty:
    st.warning(
        "No dates are currently marked Open or Closed. Gray dates are not exported."
    )
else:
    preview_df = output_df.copy()
    preview_df["*Eff"] = preview_df["*Eff"].apply(
        lambda value: (
            f"{value.month}/{value.day}/{value.year} "
            f"{value.hour}:{value.minute:02d}"
        )
    )

    st.subheader("Output preview")
    st.dataframe(
        preview_df,
        hide_index=True,
        use_container_width=True,
        height=450,
    )

    st.caption(f"The output contains {len(output_df):,} records.")

    output_name = st.text_input(
        "Output filename",
        value="Calendar_Import",
    )

    safe_output_name = re.sub(
        r"[^A-Za-z0-9_-]+",
        "_",
        output_name.strip(),
    ).strip("_")

    if not safe_output_name:
        safe_output_name = "Calendar_Import"

    if not description.strip():
        st.error("Enter a description before downloading.")
    else:
        excel_file = create_excel_file(output_df)

        st.download_button(
            label="📥 Download Calendar Import File",
            data=excel_file,
            file_name=f"{safe_output_name}.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            type="primary",
            use_container_width=True,
        )
