# -*- coding: utf-8 -*-













# APP #2: INTAKES MORE LIKE EXPECTED FILE

import io
import re
import calendar as pycalendar
from html import escape
from datetime import datetime, timedelta

import pandas as pd
import streamlit as st

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# PAGE SETUP
# ============================================================

st.set_page_config(
    page_title="Fiscal Schedule Entry",
    page_icon="📅",
    layout="wide"
)

st.title("📅 Fiscal Schedule Entry Portal")

st.markdown(
    """
    Enter operating schedules by selection group and fiscal week.

    Use the bulk-entry tools to apply the same schedule across several
    weeks, then make individual adjustments in the schedule matrix.
    """
)


# ============================================================
# CONSTANTS
# ============================================================

DAY_OPTIONS = ["Su", "M", "Tu", "W", "Th", "F", "Sa"]

DEFAULT_WEEKS = [41, 42, 43, 44, 45, 46, 47]

DEFAULT_GROUPS = [
    "DF89 4X A",
    "DF89 5X A",
    "DF89 5X B",
    "REGIONAL 4X B",
    "860",
]

SCHEDULE_PRESETS = {
    "Custom": [],
    "Monday / Tuesday / Thursday / Saturday": ["M", "Tu", "Th", "Sa"],
    "Sunday / Monday / Tuesday / Friday / Saturday": [
        "Su", "M", "Tu", "F", "Sa"
    ],
    "Sunday / Monday / Tuesday / Wednesday / Thursday / Friday / Saturday": [
        "Su", "M", "Tu", "W", "Th", "F", "Sa"
    ],
    "Sunday / Monday / Wednesday / Friday / Saturday": [
        "Su", "M", "W", "F", "Sa"
    ],
    "Monday / Tuesday / Wednesday / Friday / Saturday": [
        "M", "Tu", "W", "F", "Sa"
    ],
    "Sunday / Tuesday / Wednesday / Friday / Saturday": [
        "Su", "Tu", "W", "F", "Sa"
    ],
    "Monday / Tuesday / Thursday / Friday / Saturday": [
        "M", "Tu", "Th", "F", "Sa"
    ],
    "Monday / Tuesday / Wednesday / Thursday / Friday / Saturday": [
        "M", "Tu", "W", "Th", "F", "Sa"
    ],
    "Monday / Tuesday / Wednesday / Thursday / Friday": [
        "M", "Tu", "W", "Th", "F"
    ],
}


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def schedule_key(group_name, week_number):
    """
    Create a stable Streamlit session-state key.
    """

    clean_group = re.sub(
        r"[^A-Za-z0-9]+",
        "_",
        group_name.strip()
    ).strip("_")

    return f"schedule_{clean_group}_{week_number}"


def special_key(group_name, week_number):
    """
    Create a session-state key for TBD and N/A statuses.
    """

    clean_group = re.sub(
        r"[^A-Za-z0-9]+",
        "_",
        group_name.strip()
    ).strip("_")

    return f"special_{clean_group}_{week_number}"


def format_schedule(selected_days, special_value):
    """
    Convert selected weekdays or a special status into one cell value.
    """

    if special_value in ["TBD", "N/A"]:
        return special_value

    if not selected_days:
        return ""

    return " ".join(
        day for day in DAY_OPTIONS
        if day in selected_days
    )


def initialize_schedule_state(groups, weeks):
    """
    Ensure every group and week has session-state values.
    """

    for group_name in groups:
        for week_number in weeks:
            day_key = schedule_key(group_name, week_number)
            status_key = special_key(group_name, week_number)

            if day_key not in st.session_state:
                st.session_state[day_key] = []

            if status_key not in st.session_state:
                st.session_state[status_key] = "Schedule"


def apply_schedule_to_weeks(
    group_name,
    selected_weeks,
    selected_days,
    special_value
):
    """
    Apply a schedule to several fiscal weeks for one selection group.
    """

    for week_number in selected_weeks:
        day_key = schedule_key(group_name, week_number)
        status_key = special_key(group_name, week_number)

        if special_value == "Schedule":
            st.session_state[day_key] = selected_days.copy()
            st.session_state[status_key] = "Schedule"
        else:
            st.session_state[day_key] = []
            st.session_state[status_key] = special_value


def copy_week(groups, source_week, destination_week):
    """
    Copy every selection group's schedule from one week to another.
    """

    for group_name in groups:
        source_day_key = schedule_key(group_name, source_week)
        source_status_key = special_key(group_name, source_week)

        destination_day_key = schedule_key(
            group_name,
            destination_week
        )
        destination_status_key = special_key(
            group_name,
            destination_week
        )

        st.session_state[destination_day_key] = (
            st.session_state[source_day_key].copy()
        )

        st.session_state[destination_status_key] = (
            st.session_state[source_status_key]
        )


def build_schedule_dataframes(groups, weeks):
    """
    Build both the wide matrix and normalized detail dataframes.
    """

    matrix_rows = []
    detail_rows = []

    for group_name in groups:
        matrix_row = {
            "Selection Group": group_name
        }

        for week_number in weeks:
            days = st.session_state[
                schedule_key(group_name, week_number)
            ]

            special_value = st.session_state[
                special_key(group_name, week_number)
            ]

            schedule_text = format_schedule(
                days,
                special_value
            )

            matrix_row[f"Fiscal Week {week_number}"] = (
                schedule_text
            )

            detail_rows.append(
                {
                    "Selection Group": group_name,
                    "Fiscal Week": week_number,
                    "Schedule": schedule_text,
                    "Sunday": "Yes" if "Su" in days else "No",
                    "Monday": "Yes" if "M" in days else "No",
                    "Tuesday": "Yes" if "Tu" in days else "No",
                    "Wednesday": "Yes" if "W" in days else "No",
                    "Thursday": "Yes" if "Th" in days else "No",
                    "Friday": "Yes" if "F" in days else "No",
                    "Saturday": "Yes" if "Sa" in days else "No",
                    "Status": special_value,
                }
            )

        matrix_rows.append(matrix_row)

    matrix_df = pd.DataFrame(matrix_rows)
    detail_df = pd.DataFrame(detail_rows)

    return matrix_df, detail_df

# FOR APP #3: SPECIFIC OUTPUT FORMAT
def build_import_file(
    matrix_df,
    cal_name,
    descr,
    fiscal_week_dates
):
    """
    Converts the schedule matrix into the calendar import format.
    """

    import_rows = []

    day_lookup = {
        "M": 0,
        "Tu": 1,
        "W": 2,
        "Th": 3,
        "F": 4,
        "Sa": 5,
        "Su": 6
    }

    for _, row in matrix_df.iterrows():

        group_name = row["Selection Group"]

        for week_col in matrix_df.columns:

            if "Fiscal Week" not in week_col:
                continue

            week_number = int(
                week_col.replace(
                    "Fiscal Week ",
                    ""
                )
            )

            schedule_text = str(
                row[week_col]
            ).strip()

            if schedule_text in [
                "",
                "TBD",
                "N/A"
            ]:
                continue

            active_days = schedule_text.split()

            week_dates = fiscal_week_dates[
                week_number
            ]

            for current_date in week_dates:

                weekday_name = [
                    "M",
                    "Tu",
                    "W",
                    "Th",
                    "F",
                    "Sa",
                    "Su"
                ][current_date.weekday()]

                if weekday_name in active_days:
                    opt_value = 2
                else:
                    opt_value = 1

                import_rows.append({
                    "*Cal": cal_name,
                    "*Eff": current_date,
                    "AllocWgt": 0,
                    "AltCal": "",
                    "Avail": 0,
                    "Descr": descr,
                    "Opt": opt_value,
                    "PerWgt": 0,
                    "Repeat": ""
                })

                import_rows.append({
                    "*Cal": cal_name,
                    "*Eff": (
                        current_date
                        + timedelta(minutes=1)
                    ),
                    "AllocWgt": 0,
                    "AltCal": "",
                    "Avail": 0,
                    "Descr": descr,
                    "Opt": 1,
                    "PerWgt": 0,
                    "Repeat": ""
                })

    output_df = pd.DataFrame(
        import_rows
    )

    if not output_df.empty:

        output_df["*Eff"] = pd.to_datetime(
            output_df["*Eff"]
        )
    
        last_index = output_df.index[-1]
    
        output_df.loc[
            last_index,
            "Repeat"
        ] = "7D"
    
        midnight_rows = output_df[
            output_df["*Eff"].dt.minute == 0
        ]

        if len(midnight_rows) > 0:

            final_midnight = midnight_rows.index[-1]

            output_df.loc[
                final_midnight,
                "Opt"
            ] = 5

    return output_df

def build_fiscal_week_date_map(
    fiscal_weeks,
    first_week_start
):
    """
    Create a mapping between fiscal weeks and actual dates.
    """

    fiscal_week_dates = {}

    for week_index, week_number in enumerate(
        fiscal_weeks
    ):

        week_start = (
            first_week_start
            + timedelta(days=week_index * 7)
        )

        fiscal_week_dates[week_number] = []

        for day_offset in range(7):

            fiscal_week_dates[
                week_number
            ].append(
                week_start
                + timedelta(days=day_offset)
            )

    return fiscal_week_dates

def build_simple_calendar_records(
    selection_groups,
    fiscal_weeks,
    fiscal_week_dates
):
    """
    Build one calendar record for every selection group and date.

    Status values:
    - OPEN
    - CLOSED
    - TBD
    - N/A
    - NOT ENTERED
    """

    weekday_codes = {
        0: "M",
        1: "Tu",
        2: "W",
        3: "Th",
        4: "F",
        5: "Sa",
        6: "Su",
    }

    records_by_group = {}

    for group_name in selection_groups:
        group_records = {}

        for week_number in fiscal_weeks:
            entry_type = st.session_state[
                special_key(
                    group_name,
                    week_number
                )
            ]

            selected_days = st.session_state[
                schedule_key(
                    group_name,
                    week_number
                )
            ]

            for calendar_date in fiscal_week_dates[
                week_number
            ]:
                weekday_code = weekday_codes[
                    calendar_date.weekday()
                ]

                if entry_type == "TBD":
                    day_status = "TBD"

                elif entry_type == "N/A":
                    day_status = "N/A"

                elif not selected_days:
                    day_status = "NOT ENTERED"

                elif weekday_code in selected_days:
                    day_status = "OPEN"

                else:
                    day_status = "CLOSED"

                group_records[calendar_date] = {
                    "status": day_status,
                    "fiscal_week": week_number,
                    "weekday_code": weekday_code,
                }

        records_by_group[group_name] = group_records

    return records_by_group


def get_calendar_months(records_by_group):
    """
    Return each year and month represented in the calendar records.
    """

    months = set()

    for group_records in records_by_group.values():
        for calendar_date in group_records:
            months.add(
                (
                    calendar_date.year,
                    calendar_date.month
                )
            )

    return sorted(months)


def create_calendar_month_html(
    year,
    month,
    group_records
):
    """
    Create one simple visual calendar month as HTML.
    """

    status_colors = {
        "OPEN": {
            "background": "#D9EAD3",
            "border": "#34A853",
            "text": "#185C37",
            "label": "OPEN",
        },
        "CLOSED": {
            "background": "#F4CCCC",
            "border": "#EA4335",
            "text": "#8B1A10",
            "label": "CLOSED",
        },
        "TBD": {
            "background": "#FFF2CC",
            "border": "#F4B400",
            "text": "#7F6000",
            "label": "TBD",
        },
        "N/A": {
            "background": "#E7E6E6",
            "border": "#A6A6A6",
            "text": "#595959",
            "label": "N/A",
        },
        "NOT ENTERED": {
            "background": "#E7E6E6",
            "border": "#A6A6A6",
            "text": "#595959",
            "label": "BLANK",
        },
    }

    calendar_builder = pycalendar.Calendar(
        firstweekday=6
    )

    month_weeks = calendar_builder.monthdayscalendar(
        year,
        month
    )

    month_name = pycalendar.month_name[month]

    html_text = f"""
    <div style="
        border: 1px solid #D9D9D9;
        border-radius: 10px;
        padding: 12px;
        margin-bottom: 18px;
        background-color: white;
        box-shadow: 0 1px 4px rgba(0,0,0,0.08);
    ">
        <div style="
            font-size: 18px;
            font-weight: 700;
            text-align: center;
            margin-bottom: 10px;
            color: #1F1F1F;
        ">
            {month_name} {year}
        </div>

        <div style="
            display: grid;
            grid-template-columns: repeat(7, minmax(70px, 1fr));
            gap: 5px;
        ">
    """

    weekday_headers = [
        "Sun",
        "Mon",
        "Tue",
        "Wed",
        "Thu",
        "Fri",
        "Sat"
    ]

    for weekday_header in weekday_headers:
        html_text += f"""
        <div style="
            background-color: #1F4E78;
            color: white;
            padding: 7px 3px;
            text-align: center;
            font-size: 12px;
            font-weight: 700;
            border-radius: 4px;
        ">
            {weekday_header}
        </div>
        """

    for calendar_week in month_weeks:
        for day_number in calendar_week:
            if day_number == 0:
                html_text += """
                <div style="
                    min-height: 64px;
                    background-color: #FAFAFA;
                    border: 1px solid #EEEEEE;
                    border-radius: 5px;
                ">
                </div>
                """

                continue

            calendar_date = datetime(
                year,
                month,
                day_number
            ).date()

            if calendar_date not in group_records:
                html_text += f"""
                <div style="
                    min-height: 64px;
                    padding: 5px;
                    background-color: #FAFAFA;
                    border: 1px solid #EEEEEE;
                    border-radius: 5px;
                    color: #A6A6A6;
                ">
                    <div style="
                        font-size: 12px;
                        font-weight: 700;
                    ">
                        {day_number}
                    </div>
                </div>
                """

                continue

            record = group_records[calendar_date]
            day_status = record["status"]
            fiscal_week = record["fiscal_week"]
            style = status_colors[day_status]

            html_text += f"""
            <div style="
                min-height: 64px;
                padding: 5px;
                background-color: {style["background"]};
                border: 2px solid {style["border"]};
                border-radius: 5px;
                color: {style["text"]};
                box-sizing: border-box;
            ">
                <div style="
                    font-size: 13px;
                    font-weight: 700;
                ">
                    {day_number}
                </div>

                <div style="
                    font-size: 10px;
                    font-weight: 700;
                    margin-top: 5px;
                ">
                    {style["label"]}
                </div>

                <div style="
                    font-size: 9px;
                    margin-top: 2px;
                ">
                    FW {fiscal_week}
                </div>
            </div>
            """

    html_text += """
        </div>
    </div>
    """

    return html_text


def display_simple_calendars(
    records_by_group
):
    """
    Display one clearly labeled calendar for each selection group.
    """

    calendar_months = get_calendar_months(
        records_by_group
    )

    legend_html = """
    <div style="
        display: flex;
        flex-wrap: wrap;
        gap: 12px;
        margin-bottom: 16px;
        font-size: 13px;
    ">
        <div>
            <span style="
                display: inline-block;
                width: 14px;
                height: 14px;
                background-color: #D9EAD3;
                border: 2px solid #34A853;
                border-radius: 3px;
                vertical-align: middle;
            "></span>
            Open
        </div>

        <div>
            <span style="
                display: inline-block;
                width: 14px;
                height: 14px;
                background-color: #F4CCCC;
                border: 2px solid #EA4335;
                border-radius: 3px;
                vertical-align: middle;
            "></span>
            Closed
        </div>

        <div>
            <span style="
                display: inline-block;
                width: 14px;
                height: 14px;
                background-color: #FFF2CC;
                border: 2px solid #F4B400;
                border-radius: 3px;
                vertical-align: middle;
            "></span>
            TBD
        </div>

        <div>
            <span style="
                display: inline-block;
                width: 14px;
                height: 14px;
                background-color: #E7E6E6;
                border: 2px solid #A6A6A6;
                border-radius: 3px;
                vertical-align: middle;
            "></span>
            N/A or not entered
        </div>
    </div>
    """

    st.markdown(
        legend_html,
        unsafe_allow_html=True
    )

    for group_name, group_records in (
        records_by_group.items()
    ):
        with st.expander(
            f"📅 {group_name}",
            expanded=True
        ):
            st.subheader(group_name)

            for year, month in calendar_months:
                calendar_html = (
                    create_calendar_month_html(
                        year=year,
                        month=month,
                        group_records=group_records
                    )
                )

                st.components.v1.html(
                    calendar_html,
                    height=650,
                    scrolling=True
                )


def safe_excel_sheet_name(
    group_name,
    used_sheet_names
):
    """
    Create a valid and unique Excel worksheet name.
    """

    cleaned_name = re.sub(
        r"[\[\]\:\*\?\/\\]+",
        "_",
        group_name
    ).strip()

    if not cleaned_name:
        cleaned_name = "Selection Group"

    cleaned_name = cleaned_name[:31]

    original_name = cleaned_name
    suffix_number = 2

    while cleaned_name in used_sheet_names:
        suffix = f"_{suffix_number}"

        cleaned_name = (
            original_name[
                :31 - len(suffix)
            ]
            + suffix
        )

        suffix_number += 1

    used_sheet_names.add(cleaned_name)

    return cleaned_name


def create_simple_calendar_excel(
    records_by_group
):
    """
    Create a separate visual calendar workbook.

    Each selection group receives a clearly labeled worksheet.
    Closed dates are red and open dates are green.
    """

    workbook = Workbook()

    default_sheet = workbook.active
    workbook.remove(default_sheet)

    used_sheet_names = set()

    open_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAD3"
    )

    closed_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC"
    )

    tbd_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC"
    )

    unavailable_fill = PatternFill(
        fill_type="solid",
        fgColor="E7E6E6"
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    white_font = Font(
        color="FFFFFF",
        bold=True
    )

    title_font = Font(
        color="1F4E78",
        bold=True,
        size=16
    )

    closed_font = Font(
        color="9C0006",
        bold=True
    )

    open_font = Font(
        color="006100",
        bold=True
    )

    orange_font = Font(
        color="9C6500",
        bold=True
    )

    gray_font = Font(
        color="595959",
        bold=True
    )

    calendar_builder = pycalendar.Calendar(
        firstweekday=6
    )

    calendar_months = get_calendar_months(
        records_by_group
    )

    for group_name, group_records in (
        records_by_group.items()
    ):
        sheet_name = safe_excel_sheet_name(
            group_name,
            used_sheet_names
        )

        worksheet = workbook.create_sheet(
            title=sheet_name
        )

        worksheet.sheet_view.showGridLines = False

        worksheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=7
        )

        title_cell = worksheet.cell(
            row=1,
            column=1,
            value=f"Selection Group: {group_name}"
        )

        title_cell.font = title_font
        title_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        worksheet.row_dimensions[1].height = 28

        legend_row = 3

        legend_entries = [
            ("Open", open_fill, open_font),
            ("Closed", closed_fill, closed_font),
            ("TBD", tbd_fill, orange_font),
            (
                "N/A or not entered",
                unavailable_fill,
                gray_font
            ),
        ]

        legend_column = 1

        for legend_text, legend_fill, legend_font in (
            legend_entries
        ):
            legend_cell = worksheet.cell(
                row=legend_row,
                column=legend_column,
                value=legend_text
            )

            legend_cell.fill = legend_fill
            legend_cell.font = legend_font
            legend_cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            legend_column += 2

        current_row = 5

        for year, month in calendar_months:
            month_name = pycalendar.month_name[month]

            worksheet.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=7
            )

            month_title_cell = worksheet.cell(
                row=current_row,
                column=1,
                value=f"{month_name} {year}"
            )

            month_title_cell.fill = header_fill
            month_title_cell.font = white_font
            month_title_cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            worksheet.row_dimensions[
                current_row
            ].height = 24

            current_row += 1

            weekday_headers = [
                "Sunday",
                "Monday",
                "Tuesday",
                "Wednesday",
                "Thursday",
                "Friday",
                "Saturday"
            ]

            for column_number, weekday_header in (
                enumerate(
                    weekday_headers,
                    start=1
                )
            ):
                header_cell = worksheet.cell(
                    row=current_row,
                    column=column_number,
                    value=weekday_header
                )

                header_cell.fill = header_fill
                header_cell.font = white_font
                header_cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
                )

            current_row += 1

            month_weeks = (
                calendar_builder.monthdayscalendar(
                    year,
                    month
                )
            )

            for month_week in month_weeks:
                for column_number, day_number in (
                    enumerate(
                        month_week,
                        start=1
                    )
                ):
                    calendar_cell = worksheet.cell(
                        row=current_row,
                        column=column_number
                    )

                    calendar_cell.alignment = Alignment(
                        horizontal="center",
                        vertical="center",
                        wrap_text=True
                    )

                    if day_number == 0:
                        calendar_cell.value = ""
                        continue

                    calendar_date = datetime(
                        year,
                        month,
                        day_number
                    ).date()

                    if calendar_date not in group_records:
                        calendar_cell.value = day_number
                        calendar_cell.font = Font(
                            color="A6A6A6"
                        )

                        continue

                    record = group_records[
                        calendar_date
                    ]

                    day_status = record["status"]
                    fiscal_week = record[
                        "fiscal_week"
                    ]

                    calendar_cell.value = (
                        f"{day_number}\n"
                        f"{day_status}\n"
                        f"FW {fiscal_week}"
                    )

                    if day_status == "OPEN":
                        calendar_cell.fill = open_fill
                        calendar_cell.font = open_font

                    elif day_status == "CLOSED":
                        calendar_cell.fill = closed_fill
                        calendar_cell.font = closed_font

                    elif day_status == "TBD":
                        calendar_cell.fill = tbd_fill
                        calendar_cell.font = orange_font

                    else:
                        calendar_cell.fill = (
                            unavailable_fill
                        )

                        calendar_cell.font = gray_font

                worksheet.row_dimensions[
                    current_row
                ].height = 52

                current_row += 1

            current_row += 2

        for column_number in range(1, 8):
            column_letter = get_column_letter(
                column_number
            )

            worksheet.column_dimensions[
                column_letter
            ].width = 16

        worksheet.freeze_panes = "A5"

    output = io.BytesIO()

    workbook.save(output)

    output.seek(0)

    return output

# FOR APP #2 - NOT SPECIFYING OUR OUTPUT FORMAT
def create_excel_workbook(matrix_df, detail_df):
    """
    Create a formatted Excel workbook in memory.
    """

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        matrix_df.to_excel(
            writer,
            index=False,
            sheet_name="Schedule Matrix"
        )

        detail_df.to_excel(
            writer,
            index=False,
            sheet_name="Schedule Detail"
        )

        workbook = writer.book

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        header_font = Font(
            color="FFFFFF",
            bold=True
        )

        selection_fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        tbd_fill = PatternFill(
            fill_type="solid",
            fgColor="FFF2CC"
        )

        unavailable_fill = PatternFill(
            fill_type="solid",
            fgColor="E7E6E6"
        )

        for worksheet in workbook.worksheets:

            worksheet.freeze_panes = "B2"
            worksheet.auto_filter.ref = worksheet.dimensions

            for header_cell in worksheet[1]:
                header_cell.fill = header_fill
                header_cell.font = header_font
                header_cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True
            )

            for row in worksheet.iter_rows(
                min_row=2,
                max_row=worksheet.max_row
            ):
                for cell in row:
                    cell.alignment = Alignment(
                        vertical="center",
                        wrap_text=True
                    )

                    if cell.column == 1:
                        cell.fill = selection_fill
                        cell.font = Font(bold=True)

                    if cell.value == "TBD":
                        cell.fill = tbd_fill

                    if cell.value == "N/A":
                        cell.fill = unavailable_fill

            for column_cells in worksheet.columns:
                max_length = 0

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:
                    cell_value = (
                        "" if cell.value is None
                        else str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        len(cell_value)
                    )

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max(max_length + 2, 12),
                    35
                )

            worksheet.row_dimensions[1].height = 30

    output.seek(0)
    return output


# ============================================================
# FISCAL WEEK CONFIGURATION
# ============================================================

st.subheader("1. Fiscal Week Setup")

week_col1, week_col2 = st.columns([2, 3])

with week_col1:
    starting_week = st.number_input(
        "Starting fiscal week",
        min_value=1,
        max_value=53,
        value=41,
        step=1
    )

with week_col2:
    number_of_weeks = st.slider(
        "Number of fiscal weeks",
        min_value=1,
        max_value=13,
        value=7
    )

fiscal_weeks = list(
    range(
        int(starting_week),
        int(starting_week) + int(number_of_weeks)
    )
)

st.caption(
    "Active fiscal weeks: "
    + ", ".join(str(week) for week in fiscal_weeks)
)

st.subheader("Week Calendar Mapping")

first_fiscal_week_start = st.date_input(
    f"Start Date of Fiscal Week {fiscal_weeks[0]}",
    value=datetime(2026, 10, 17).date()
)

fiscal_week_dates = build_fiscal_week_date_map(
    fiscal_weeks,
    first_fiscal_week_start
)

preview_rows = []

for week_number in fiscal_weeks:

    preview_rows.append({
        "Fiscal Week": week_number,
        "Start Date": fiscal_week_dates[
            week_number
        ][0],
        "End Date": fiscal_week_dates[
            week_number
        ][6]
    })

st.dataframe(
    pd.DataFrame(preview_rows),
    hide_index=True,
    use_container_width=True
)


# ============================================================
# SELECTION GROUP SETUP
# ============================================================

st.subheader("2. Selection Groups")

if "selection_groups" not in st.session_state:
    st.session_state.selection_groups = DEFAULT_GROUPS.copy()

group_text = st.text_area(
    "Enter one selection group per line",
    value="\n".join(st.session_state.selection_groups),
    height=150,
    help=(
        "Add, remove, or rename selection groups. "
        "Blank lines are ignored."
    )
)

selection_groups = [
    group.strip()
    for group in group_text.splitlines()
    if group.strip()
]

# Remove duplicate names while preserving order
selection_groups = list(dict.fromkeys(selection_groups))

st.session_state.selection_groups = selection_groups

if not selection_groups:
    st.warning(
        "Enter at least one selection group to continue."
    )
    st.stop()

initialize_schedule_state(
    selection_groups,
    fiscal_weeks
)


# ============================================================
# BULK ENTRY TOOLS
# ============================================================

st.subheader("3. Quick Schedule Entry")

st.markdown(
    """
    Pick a selection group, a schedule, and the fiscal weeks that
    should receive that schedule.
    """
)

bulk_col1, bulk_col2 = st.columns(2)

with bulk_col1:
    bulk_group = st.selectbox(
        "Selection group",
        options=selection_groups
    )

    preset_name = st.selectbox(
        "Schedule template",
        options=list(SCHEDULE_PRESETS.keys())
    )

with bulk_col2:
    bulk_weeks = st.multiselect(
        "Apply to fiscal weeks",
        options=fiscal_weeks,
        default=fiscal_weeks
    )

    bulk_status = st.radio(
        "Entry type",
        options=["Schedule", "TBD", "N/A"],
        horizontal=True
    )

preset_days = SCHEDULE_PRESETS[preset_name]

if preset_name == "Custom":
    bulk_days = st.multiselect(
        "Select operating days",
        options=DAY_OPTIONS,
        default=[],
        key="bulk_custom_days"
    )
else:
    bulk_days = st.multiselect(
        "Operating days",
        options=DAY_OPTIONS,
        default=preset_days,
        key=f"bulk_days_{preset_name}"
    )

if bulk_status != "Schedule":
    st.info(
        f"The selected weeks will be marked as {bulk_status}."
    )

if st.button(
    "Apply Schedule to Selected Weeks",
    type="primary",
    use_container_width=True
):
    if not bulk_weeks:
        st.error(
            "Select at least one fiscal week."
        )
    elif (
        bulk_status == "Schedule"
        and not bulk_days
    ):
        st.error(
            "Select at least one operating day, "
            "or choose TBD or N/A."
        )
    else:
        apply_schedule_to_weeks(
            group_name=bulk_group,
            selected_weeks=bulk_weeks,
            selected_days=bulk_days,
            special_value=bulk_status
        )

        st.success(
            f"Updated {bulk_group} for "
            f"{len(bulk_weeks)} fiscal week(s)."
        )

        st.rerun()


# ============================================================
# COPY WEEK TOOL
# ============================================================

with st.expander("Copy all schedules from one week to another"):

    copy_col1, copy_col2, copy_col3 = st.columns(
        [2, 2, 1]
    )

    with copy_col1:
        source_week = st.selectbox(
            "Copy from fiscal week",
            options=fiscal_weeks,
            key="source_week"
        )

    with copy_col2:
        destination_options = [
            week for week in fiscal_weeks
            if week != source_week
        ]

        destination_week = st.selectbox(
            "Copy into fiscal week",
            options=destination_options,
            key="destination_week"
        )

    with copy_col3:
        st.write("")
        st.write("")

        if st.button(
            "Copy Week",
            use_container_width=True
        ):
            copy_week(
                selection_groups,
                source_week,
                destination_week
            )

            st.success(
                f"Fiscal week {source_week} was copied "
                f"into fiscal week {destination_week}."
            )

            st.rerun()


# ============================================================
# DETAILED SCHEDULE MATRIX
# ============================================================

st.subheader("4. Review and Adjust Schedules")

st.caption(
    "Expand a selection group to make individual week-by-week changes."
)

for group_name in selection_groups:

    with st.expander(
        f"📂 {group_name}",
        expanded=False
    ):

        for week_number in fiscal_weeks:

            st.markdown(
                f"**Fiscal Week {week_number}**"
            )

            schedule_col, status_col = st.columns(
                [4, 2]
            )

            with status_col:
                entry_type = st.selectbox(
                    "Entry type",
                    options=["Schedule", "TBD", "N/A"],
                    key=special_key(
                        group_name,
                        week_number
                    ),
                    label_visibility="collapsed"
                )

            with schedule_col:
                selected_days = st.multiselect(
                    f"Operating days for week {week_number}",
                    options=DAY_OPTIONS,
                    key=schedule_key(
                        group_name,
                        week_number
                    ),
                    disabled=entry_type != "Schedule",
                    label_visibility="collapsed",
                    placeholder="Select operating days"
                )

            if entry_type != "Schedule":
                st.caption(
                    f"Fiscal week {week_number} is marked "
                    f"{entry_type}."
                )

            st.divider()


# ============================================================
# PREVIEW
# ============================================================

st.subheader("5. Schedule Preview")

matrix_df, detail_df = build_schedule_dataframes(
    selection_groups,
    fiscal_weeks
)

st.dataframe(
    matrix_df,
    hide_index=True,
    use_container_width=True
)

blank_entries = 0

for group_name in selection_groups:
    for week_number in fiscal_weeks:

        day_values = st.session_state[
            schedule_key(group_name, week_number)
        ]

        special_value = st.session_state[
            special_key(group_name, week_number)
        ]

        if (
            special_value == "Schedule"
            and not day_values
        ):
            blank_entries += 1

if blank_entries:
    st.warning(
        f"{blank_entries} schedule cell(s) are still blank. "
        "You may continue, but review them before downloading."
    )
else:
    st.success(
        "All schedule cells contain a schedule, TBD, or N/A."
    )

# ============================================================
# SIMPLE VISUAL CALENDAR
# ============================================================

st.subheader("7. Simple Calendar View")

st.markdown(
    """
    Each selection group is displayed separately.

    Closed dates are shown in red. Open dates are shown in green.
    Weeks marked TBD are yellow, and unavailable or incomplete
    schedules are gray.
    """
)

simple_calendar_records = (
    build_simple_calendar_records(
        selection_groups=selection_groups,
        fiscal_weeks=fiscal_weeks,
        fiscal_week_dates=fiscal_week_dates
    )
)

display_simple_calendars(
    simple_calendar_records
)

# ============================================================
# EXCEL DOWNLOAD
# ============================================================

st.subheader("8. Import File Settings")

cal_name = st.text_input(
    "*Cal Value",
    value="SHIP_DF_860_TG"
)

descr = st.text_input(
    "Description",
    value="turkeys 2026"
)

st.subheader("7. Download Excel File")

submission_name = st.text_input(
    "Submission name",
    value=(
        "Fiscal_Schedule_"
        + datetime.now().strftime("%Y%m%d")
    )
)

safe_filename = re.sub(
    r"[^A-Za-z0-9_-]+",
    "_",
    submission_name.strip()
).strip("_")

if not safe_filename:
    safe_filename = "Fiscal_Schedule"

# FOR APP #2

# excel_file = create_excel_workbook(
#     matrix_df,
#     detail_df
# )

import_df = build_import_file(
    matrix_df=matrix_df,
    cal_name=cal_name,
    descr=descr,
    fiscal_week_dates=fiscal_week_dates
)

excel_buffer = io.BytesIO()

with pd.ExcelWriter(
    excel_buffer,
    engine="openpyxl"
) as writer:

    import_df.to_excel(
        writer,
        sheet_name="Calendar Import",
        index=False
    )

excel_buffer.seek(0)

excel_file = excel_buffer

st.download_button(
    label="📥 Download Completed Schedule",
    data=excel_file,
    file_name=f"{safe_filename}.xlsx",
    mime=(
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),
    type="primary",
    use_container_width=True
)

st.markdown("#### Simple Calendar Version")

st.caption(
    "This workbook contains a visual calendar worksheet "
    "for each selection group. Closed dates are highlighted red."
)

simple_excel_file = create_simple_calendar_excel(
    simple_calendar_records
)

simple_file_name = (
    f"{safe_filename}_Simple_Version.xlsx"
)

st.download_button(
    label="📥 Download Simple Calendar Version",
    data=simple_excel_file,
    file_name=simple_file_name,
    mime=(
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet"
    ),
    type="secondary",
    use_container_width=True,
    key="download_simple_calendar"
)      

# WORKING APP #1: BASE SUGGESTION OF INPUTTING AND OUTPUTTING CALENDAR FILES

# import streamlit as st
# import pandas as pd
# from datetime import datetime, timedelta
# import io

# # --------------------------------------------------
# # Page Setup
# # --------------------------------------------------

# st.set_page_config(
#     page_title="Schedule Form Submission",
#     page_icon="📅",
#     layout="centered"
# )

# st.title("📅 Facility Schedule Submission Portal")

# st.markdown("""
# Please populate the operational windows below.

# The visual calendar dynamically updates to reflect your inputs.

# Once submitted, an Excel file will be generated and available for download.
# """)

# # --------------------------------------------------
# # Inputs
# # --------------------------------------------------

# st.subheader("1. Parameters & Configuration")

# facility_name = st.text_input(
#     "What is the exact name of the Location or Facility?",
#     placeholder="Regional Distribution Center"
# )

# status = st.selectbox(
#     "What structural classification applies to this date span?",
#     [
#         "Normal Operations",
#         "Seasonal Closure",
#         "Maintenance Close"
#     ]
# )

# col1, col2 = st.columns(2)

# with col1:
#     open_date = st.date_input(
#         "Operational Window START Date",
#         datetime.today().date()
#     )

# with col2:
#     close_date = st.date_input(
#         "Operational Window END Date",
#         datetime.today().date() + timedelta(days=5)
#     )

# notes = st.text_area(
#     "Provide any contextual notes",
#     placeholder="Optional..."
# )

# # --------------------------------------------------
# # Timeline Calculation
# # --------------------------------------------------

# st.subheader("2. Real-Time Dynamic Timeline View")

# window_days = []

# curr = open_date

# while curr <= close_date:
#     window_days.append(curr)
#     curr += timedelta(days=1)

# is_closure_type = (
#     "Close" in status
#     or "Closure" in status
# )

# bg_color = "#EA4335" if is_closure_type else "#34A853"
# text_status = "CLOSED" if is_closure_type else "OPEN"

# calendar_html = """
# <div style='display:flex;
#             flex-wrap:wrap;
#             gap:6px;
#             padding:5px 0;'>
# """

# for day in window_days[:31]:

#     calendar_html += f"""
#     <div style="
#         background-color:{bg_color};
#         color:white;
#         padding:6px;
#         border-radius:4px;
#         text-align:center;
#         width:55px;
#         font-size:11px;
#         box-shadow:1px 1px 3px rgba(0,0,0,.1);">

#         {day.strftime('%b %d')}
#         <br>
#         <b>{text_status}</b>
#     </div>
#     """

# if len(window_days) > 31:
#     calendar_html += """
#     <div style="
#         align-self:center;
#         padding-left:10px;
#         font-size:12px;">
#         + more days
#     </div>
#     """

# calendar_html += "</div>"

# st.components.v1.html(
#     calendar_html,
#     height=90,
#     scrolling=False
# )

# # --------------------------------------------------
# # Submit
# # --------------------------------------------------

# if st.button(
#     "🚀 Generate Schedule File",
#     use_container_width=True
# ):

#     if not facility_name.strip():

#         st.error(
#             "Please provide a Facility Name."
#         )

#     elif open_date > close_date:

#         st.error(
#             "Start date cannot be after End date."
#         )

#     else:

#         submission_data = {
#             "Facility Location": [facility_name],
#             "Schedule Designation": [status],
#             "Start Effective Date": [open_date],
#             "End Terminating Date": [close_date],
#             "Total Span Days": [len(window_days)],
#             "Notes": [
#                 notes if notes else "None Provided"
#             ]
#         }

#         df = pd.DataFrame(submission_data)

#         excel_buffer = io.BytesIO()

#         with pd.ExcelWriter(
#             excel_buffer,
#             engine="openpyxl"
#         ) as writer:

#             df.to_excel(
#                 writer,
#                 index=False,
#                 sheet_name="Submission"
#             )

#         excel_buffer.seek(0)

#         st.success(
#             "Excel file generated successfully."
#         )

#         st.download_button(
#             label="📥 Download Excel File",
#             data=excel_buffer,
#             file_name=f"{facility_name.replace(' ','_')}_Schedule.xlsx",
#             mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
#         )